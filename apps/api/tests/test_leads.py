import asyncio
import io
from datetime import UTC, datetime, timedelta

import pytest
from conftest import API, create_user, login, new_client
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import func, select

from app.core.config import settings
from app.core.db import SessionLocal
from app.models import AuditLog, Lead, LeadEvent, LeadStatus, Role
from app.services import telegram

PUBLIC = "/api/v1/public/leads"


def lead_body(**kw):
    return {"full_name": "Ali Valiyev", "phone": "90 123 45 67", "consent": True, "elapsed_ms": 5000, **kw}


async def lead_count() -> int:
    async with SessionLocal() as s:
        return await s.scalar(select(func.count()).select_from(Lead)) or 0


async def make_lead(**kw) -> int:
    async with SessionLocal() as s:
        lead = Lead(
            full_name=kw.pop("full_name", "Test Bemor"),
            phone=kw.pop("phone", "+998901112233"),
            consent=True,
            **kw,
        )
        s.add(lead)
        await s.flush()
        s.add(LeadEvent(lead_id=lead.id, kind="created"))
        await s.commit()
        return lead.id


# ───────────────────────── Public ariza ─────────────────────────


async def test_public_lead_created_and_normalized(client):
    res = await client.post(PUBLIC, json=lead_body(message="Bosh og‘riyapti", page="/doctors/x"))
    assert res.status_code == 201 and res.json() == {"ok": True}
    async with SessionLocal() as s:
        lead = await s.scalar(select(Lead))
        assert lead.phone == "+998901234567" and lead.consent and lead.status == LeadStatus.new


async def test_honeypot_and_too_fast_not_saved(client):
    assert (await client.post(PUBLIC, json=lead_body(website="http://spam"))).status_code == 201
    assert (await client.post(PUBLIC, json=lead_body(elapsed_ms=300))).status_code == 201
    assert await lead_count() == 0


@pytest.mark.parametrize("phone", ["12345", "+7 999 123 45 67", "abcdefghij", "90 123 45 6"])
async def test_phone_validation(client, phone):
    res = await client.post(PUBLIC, json=lead_body(phone=phone))
    assert res.status_code == 422
    assert await lead_count() == 0


async def test_public_lead_validation(client):
    assert (await client.post(PUBLIC, json=lead_body(consent=False))).status_code == 422
    assert (await client.post(PUBLIC, json=lead_body(extra="x"))).status_code == 422
    assert (await client.post(PUBLIC, json=lead_body(full_name="A" * 121))).status_code == 422
    assert (await client.post(PUBLIC, json=lead_body(page="javascript:alert(1)"))).status_code == 422


async def test_dedup_same_phone(client):
    await client.post(PUBLIC, json=lead_body(message="birinchi"))
    await client.post(PUBLIC, json=lead_body(phone="+998 90 123 45 67", message="ikkinchi"))
    assert await lead_count() == 1
    async with SessionLocal() as s:
        lead = await s.scalar(select(Lead))
        assert "birinchi" in lead.message and "ikkinchi" in lead.message
        kinds = (await s.scalars(select(LeadEvent.kind).where(LeadEvent.lead_id == lead.id))).all()
        assert "duplicate" in kinds


async def test_public_lead_rate_limit(client):
    headers = {"x-real-ip": "203.0.113.7"}
    for i in range(5):
        res = await client.post(PUBLIC, json=lead_body(phone=f"90 000 00 0{i}"), headers=headers)
        assert res.status_code == 201
    res = await client.post(PUBLIC, json=lead_body(phone="90 000 00 09"), headers=headers)
    assert res.status_code == 429


# ───────────────────────── Telegram webhook ─────────────────────────


@pytest.fixture
def tg(monkeypatch):
    calls: list[tuple[str, dict]] = []

    async def fake_call(method: str, payload: dict) -> dict:
        calls.append((method, payload))
        return {"message_id": 777}

    monkeypatch.setattr(settings, "telegram_bot_token", SecretStr("123:abc"))
    monkeypatch.setattr(settings, "telegram_chat_id", "-100500")
    monkeypatch.setattr(settings, "telegram_webhook_secret", SecretStr("hook-secret-123"))
    monkeypatch.setattr(telegram, "call", fake_call)
    return calls


def callback(lead_id: int, user_id: int, action: str = "take", chat: str = "-100500") -> dict:
    return {
        "update_id": user_id,
        "callback_query": {
            "id": f"cb{user_id}",
            "from": {"id": user_id, "first_name": f"Admin{user_id}"},
            "message": {"chat": {"id": int(chat)}, "message_id": 777},
            "data": f"{action}:{lead_id}",
        },
    }


async def test_webhook_wrong_secret(client, tg):
    lead_id = await make_lead()
    res = await client.post("/api/v1/telegram/webhook", json=callback(lead_id, 1))
    assert res.status_code == 403
    res = await client.post(
        "/api/v1/telegram/webhook", json=callback(lead_id, 1), headers={"x-telegram-bot-api-secret-token": "wrong"}
    )
    assert res.status_code == 403


async def test_webhook_take_race(client, tg):
    lead_id = await make_lead(tg_message_id=777)
    h = {"x-telegram-bot-api-secret-token": "hook-secret-123"}
    r1, r2 = await asyncio.gather(
        client.post("/api/v1/telegram/webhook", json=callback(lead_id, 1), headers=h),
        client.post("/api/v1/telegram/webhook", json=callback(lead_id, 2), headers=h),
    )
    assert r1.status_code == r2.status_code == 200
    async with SessionLocal() as s:
        lead = await s.get(Lead, lead_id)
        assert lead.status == LeadStatus.in_progress
        assert lead.taken_by in ("Admin1", "Admin2")
        taken = (await s.scalars(select(LeadEvent).where(LeadEvent.lead_id == lead_id, LeadEvent.kind == "taken"))).all()
        assert len(taken) == 1
    answers = [p["text"] for m, p in tg if m == "answerCallbackQuery"]
    assert sum(a == "Ariza sizga biriktirildi" for a in answers) == 1
    assert any("olgan" in a for a in answers)


async def test_webhook_foreign_chat_ignored(client, tg):
    lead_id = await make_lead()
    h = {"x-telegram-bot-api-secret-token": "hook-secret-123"}
    res = await client.post("/api/v1/telegram/webhook", json=callback(lead_id, 1, chat="-999"), headers=h)
    assert res.status_code == 200
    async with SessionLocal() as s:
        assert (await s.get(Lead, lead_id)).status == LeadStatus.new


# ───────────────────────── Admin: arizalar ─────────────────────────


async def test_admin_take_race(operator_client):
    c, _ = operator_client
    lead_id = await make_lead()
    other = await create_user(Role.admin)
    async with new_client() as c2:
        await login(c2, other)
        r1, r2 = await asyncio.gather(c.post(f"{API}/leads/{lead_id}/take"), c2.post(f"{API}/leads/{lead_id}/take"))
    codes = sorted([r1.status_code, r2.status_code])
    assert codes == [200, 409]
    loser = r1 if r1.status_code == 409 else r2
    assert loser.json()["detail"] == "already_taken"
    winner = r1 if r1.status_code == 200 else r2
    assert winner.json()["status"] == "in_progress"
    assert [e["kind"] for e in winner.json()["events"]] == ["created", "taken"]


async def test_admin_take_conflicts_with_telegram(operator_client, tg):
    c, _ = operator_client
    lead_id = await make_lead(tg_message_id=777)
    h = {"x-telegram-bot-api-secret-token": "hook-secret-123"}
    assert (await c.post(f"{API}/leads/{lead_id}/take")).status_code == 200
    # Telegram xabari sinxronlandi (editMessageText)
    assert any(m == "editMessageText" for m, _ in tg)
    await c.post("/api/v1/telegram/webhook", json=callback(lead_id, 5), headers=h)
    async with SessionLocal() as s:
        lead = await s.get(Lead, lead_id)
        assert lead.taken_by != "Admin5"


async def test_admin_patch_status_and_note(operator_client, monkeypatch):
    c, op = operator_client
    lead_id = await make_lead(tg_message_id=1)

    async def boom(lead):
        raise RuntimeError("telegram down")

    monkeypatch.setattr(telegram, "sync", boom)
    res = await c.patch(f"{API}/leads/{lead_id}", json={"status": "booked", "admin_note": "Ertaga 10:00"})
    assert res.status_code == 200, res.text  # Telegram xatosi API'ni buzmaydi
    body = res.json()
    assert body["status"] == "booked" and body["admin_note"] == "Ertaga 10:00"
    assert body["taken_by"] == op.name
    kinds = [e["kind"] for e in body["events"]]
    assert "status" in kinds and "note" in kinds
    assert (await c.patch(f"{API}/leads/{lead_id}", json={"status": "nope"})).status_code == 422
    assert (await c.patch(f"{API}/leads/{lead_id}", json={"status": None})).status_code == 422
    assert (await c.patch(f"{API}/leads/999999", json={"status": "spam"})).status_code == 404


async def test_leads_list_filters_and_counts(operator_client):
    c, _ = operator_client
    await make_lead(full_name="Anvar Karimov", phone="+998901234567")
    await make_lead(full_name="Bobur 100%_test", phone="+998935550011", status=LeadStatus.booked)
    await make_lead(full_name="Dilnoza", phone="+998977770000", status=LeadStatus.spam)

    res = (await c.get(f"{API}/leads")).json()
    assert res["total"] == 3 and res["counts"]["all"] == 3 and res["counts"]["booked"] == 1
    assert {"id", "full_name", "phone", "status", "tg_delivery", "created_at"} <= set(res["items"][0])

    res = (await c.get(f"{API}/leads", params={"status": "booked"})).json()
    assert res["total"] == 1 and res["counts"]["all"] == 3  # counts status filtrini hisobga olmaydi

    assert (await c.get(f"{API}/leads", params={"q": "anvar"})).json()["total"] == 1
    assert (await c.get(f"{API}/leads", params={"q": "93 555"})).json()["total"] == 1
    # LIKE maxsus belgilari escape qilinadi: "%" faqat literal "%" bor ismga mos (hammasiga emas)
    assert (await c.get(f"{API}/leads", params={"q": "%"})).json()["total"] == 1
    assert (await c.get(f"{API}/leads", params={"q": "100%_"})).json()["total"] == 1

    today = datetime.now(UTC).date()
    res = (await c.get(f"{API}/leads", params={"date_from": (today + timedelta(days=2)).isoformat()})).json()
    assert res["total"] == 0 and res["counts"]["all"] == 0
    assert (await c.get(f"{API}/leads", params={"limit": 101})).status_code == 422

    first = (await c.get(f"{API}/leads")).json()["items"][0]
    detail = (await c.get(f"{API}/leads/{first['id']}")).json()
    assert detail["consent"] is True and detail["events"][0]["kind"] == "created"


async def test_dashboard_numbers(admin_client):
    c, _ = admin_client
    await make_lead()
    await make_lead(phone="+998900000001", status=LeadStatus.booked)
    d = (await c.get(f"{API}/dashboard")).json()
    assert d["today"]["leads"] == 2
    assert d["waiting"]["count"] == 1 and d["waiting"]["oldest_minutes"] is not None
    assert d["booked_14d"] == {"booked": 1, "total": 2}
    assert len(d["chart"]) == 14 and d["chart"][-1]["total"] == 2
    assert len(d["waiting_list"]) == 1
    assert d["telegram"]["enabled"] is False
    assert d["content_todo"]["doctors_without_photo"] == 0


# ───────────────────────── Eksport ─────────────────────────


async def test_export_xlsx_escapes_formulas(admin_client):
    c, _ = admin_client
    await make_lead(full_name="=HYPERLINK(\"http://evil\",\"x\")", message="+cmd|' /C calc'!A0")
    await make_lead(full_name="@SUM(1,2)", phone="+998900000002", admin_note="-2+3")
    await make_lead(full_name="Oddiy Ism", phone="+998900000003")

    res = await c.get(f"{API}/leads/export")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "ID" and len(rows) == 4
    for row in rows[1:]:
        for value in row:
            if isinstance(value, str) and value:
                assert value[0] not in "=+-@", value
    names = {r[2] for r in rows[1:]}
    assert "'=HYPERLINK(\"http://evil\",\"x\")" in names and "'@SUM(1,2)" in names and "Oddiy Ism" in names
    # Hech bir katak formula sifatida saqlanmagan
    assert all(cell.data_type != "f" for row in ws.iter_rows() for cell in row)

    async with SessionLocal() as s:
        log = await s.scalar(select(AuditLog).where(AuditLog.action == "leads_export"))
        assert log and log.meta["rows"] == 3 and log.ip is not None


async def test_export_admin_only(operator_client):
    c, _ = operator_client
    assert (await c.get(f"{API}/leads/export")).status_code == 403
